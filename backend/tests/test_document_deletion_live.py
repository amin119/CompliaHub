"""Live-infra tests for Phase 8's per-document deletion helpers — real
Qdrant and real MinIO, both safe to exercise against the actual shared
project instances since neither has any cross-document risk (confirmed:
Qdrant points and MinIO objects are never shared across documents, unlike
Neo4j's entity nodes).

Deliberately does NOT call the full `DELETE /documents/{id}` endpoint or
`document_deletion.delete_document`/`document_rebuild
.rebuild_graph_from_remaining_documents` against the real shared corpus —
this project's shared dev database has accumulated 311 documents from
years of test runs (a known, documented "test pollution" limitation), and
the full-rebuild strategy (the user's own resolved decision, see
docs/phase-8-scaling.md) reprocesses every one of them on any single
delete. A real run would take ~110+ minutes (Voyage's own ~21s inter-call
pacing alone) and burn real API spend against production-adjacent shared
state — confirmed with the user via `AskUserQuestion` to skip this in favor
of `graph_store.clear_entity_graph`'s Cypher being separately verified
live against a temporary, isolated Neo4j container instead (see
docs/phase-8-scaling.md's verification section for that result). The full
endpoint's orchestration logic (call order, which functions get invoked)
is covered by `test_document_deletion.py`'s mocked unit tests instead.
"""

import io
import uuid

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.core.db import engine
from app.main import app
from app.services import storage, vector_store
from app.tasks.celery_app import celery_app


def _infra_available() -> bool:
    try:
        with engine.connect() as conn:
            conn.execute(text("SELECT 1 FROM documents LIMIT 1"))
        vector_store.get_qdrant_client().get_collections()
        storage.get_minio_client().list_buckets()
        return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(
    not _infra_available(),
    reason="requires docker compose up (postgres+redis+qdrant+minio) with migrations applied",
)


@pytest.fixture(autouse=True, scope="module")
def _eager_celery():
    celery_app.conf.task_always_eager = True
    celery_app.conf.task_eager_propagates = True
    yield
    celery_app.conf.task_always_eager = False
    celery_app.conf.task_eager_propagates = False


@pytest.fixture(autouse=True)
def _mock_embed_stage(monkeypatch):
    """Same convention as `test_documents_api.py` — this test exercises
    real Qdrant deletion, not real Voyage embedding calls.
    """
    from app.services import embedding

    def _fake_embed_texts(texts, input_type):
        return [[0.0] * embedding.EMBEDDING_DIM for _ in texts]

    monkeypatch.setattr(embedding, "embed_texts", _fake_embed_texts)


def _sample_docx_bytes(nonce: str) -> bytes:
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Phase 8 Deletion Test Standard", level=1)
    doc.add_heading(f"A.{nonce} Test Clause", level=2)
    doc.add_paragraph(f"This is a test clause body {nonce}.")
    doc.save(buf)
    return buf.getvalue()


def test_delete_by_document_removes_only_that_documents_qdrant_points():
    client = TestClient(app)
    nonce = uuid.uuid4().hex[:8]
    data = _sample_docx_bytes(nonce)

    response = client.post(
        "/documents",
        files={"file": (f"deletion_test_{nonce}.docx", data, "application/octet-stream")},
    )
    document_id = uuid.UUID(response.json()["id"])
    assert client.get(f"/documents/{document_id}").json()["status"] == "ready"

    qdrant_client = vector_store.get_qdrant_client()
    assert vector_store.count_by_document(qdrant_client, document_id) >= 1

    vector_store.delete_by_document(qdrant_client, document_id)

    assert vector_store.count_by_document(qdrant_client, document_id) == 0


def test_storage_delete_document_removes_the_real_object():
    minio_client = storage.get_minio_client()
    object_key = f"phase8-deletion-test/{uuid.uuid4().hex}.docx"
    storage.upload_document(minio_client, object_key, b"test content", "application/octet-stream")

    # Confirms the object really exists first — otherwise a bug that made
    # `delete_document` a silent no-op would pass this test vacuously.
    minio_client.stat_object(storage.DOCUMENTS_BUCKET, object_key)

    storage.delete_document(minio_client, object_key)

    with pytest.raises(Exception):
        minio_client.stat_object(storage.DOCUMENTS_BUCKET, object_key)
