"""Platform Phase 8: a one-off load test proving (or refuting) Phase 3's own
claim that LLM extraction is the ingestion pipeline's real bottleneck — NOT
a permanent feature, same "tune against real use, don't over-engineer"
spirit as `generate_eval_questions.py`.

Uploads a synthetic corpus through the REAL `POST /documents` API (the real
HTTP+Celery+Docling path, not an in-process shortcut), then reports real
per-stage timing pulled from `ProcessingJob.started_at`/`finished_at`
(already exists per document/stage — no new instrumentation needed).

**Honest, deliberately reduced scope**: the original plan assumed a
~100-300 synthetic document run. Before running this script, a live check
of the real shared dev database found it already holds 311 documents
accumulated from years of test runs (this project's own long-documented
"shared dev/test database pollution" limitation) — directly relevant here
for the first time, since Phase 8's "full rebuild on delete" strategy makes
total corpus size a real operational cost (a real document delete now
reprocesses every document in the table). Adding another 100-300 synthetic
documents on top of that would make the exact problem this phase's own
verification just surfaced meaningfully worse for no added statistical
value. This script instead runs a smaller batch (default 20) — enough for a
real, stable per-stage timing distribution — and the resulting report
EXTRAPOLATES throughput/bottleneck claims mathematically from those real
per-document numbers to larger hypothetical corpus sizes, rather than
literally proving it by uploading thousands of files.

Usage (from `backend/`, with the real stack up via `docker compose up -d`
and a FastAPI process running — this project runs the API natively on the
host, not in Docker, so start one first, e.g.
`uv run uvicorn app.main:app --host 0.0.0.0 --port 8010`):
    uv run python scripts/load_test_ingestion.py --api-url http://localhost:8010 [--count 20]
"""

import argparse
import io
import statistics
import sys
import time
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import truststore  # noqa: E402

truststore.inject_into_ssl()

import httpx  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.core.db import SessionLocal  # noqa: E402
from app.models.document import ProcessingJob  # noqa: E402

_STAGE_ORDER = ["parse", "chunk", "embed", "extract", "resolve_and_load"]


def _generate_docx(nonce: str, clause_count: int = 5) -> bytes:
    """A realistic multi-chunk document, not one giant blob — clause-numbered
    headings so the existing chunker produces `clause_count` real chunks,
    same technique `tests/test_documents_api.py` already uses for
    `_sample_docx_bytes`.
    """
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading(f"Load Test Standard {nonce}", level=1)
    for i in range(1, clause_count + 1):
        doc.add_heading(f"A.{i} Load test clause {i}", level=2)
        doc.add_paragraph(
            f"This is synthetic clause body {i} for load-test document {nonce}. "
            "It exists purely to give the chunker and downstream pipeline "
            "stages realistic, non-trivial text to process."
        )
    doc.save(buf)
    return buf.getvalue()


def _upload_and_wait(
    client: httpx.Client, api_url: str, data: bytes, filename: str, timeout_s: float
) -> str:
    response = client.post(
        f"{api_url}/documents", files={"file": (filename, data, "application/octet-stream")}
    )
    response.raise_for_status()
    document_id = response.json()["id"]

    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        status = client.get(f"{api_url}/documents/{document_id}").json()
        if status["status"] in ("ready", "failed"):
            return document_id
        time.sleep(1)
    raise TimeoutError(f"document {document_id} did not finish within {timeout_s}s")


def _stage_durations(db, document_ids: list[str]) -> dict[str, list[float]]:
    durations: dict[str, list[float]] = {stage: [] for stage in _STAGE_ORDER}
    jobs = db.scalars(
        select(ProcessingJob).where(ProcessingJob.document_id.in_(document_ids))
    )
    for job in jobs:
        if job.task_name in durations and job.started_at and job.finished_at:
            durations[job.task_name].append((job.finished_at - job.started_at).total_seconds())
    return durations


def _report(durations: dict[str, list[float]], document_count: int) -> str:
    lines = ["# Phase 8 Load Test Results\n", f"Documents ingested: {document_count}\n"]
    lines.append("| Stage | n | mean (s) | p50 (s) | p95 (s) | total (s) |")
    lines.append("|---|---|---|---|---|---|")

    total_by_stage = {}
    for stage in _STAGE_ORDER:
        values = durations[stage]
        if not values:
            lines.append(f"| {stage} | 0 | - | - | - | - |")
            continue
        values_sorted = sorted(values)
        mean = statistics.mean(values)
        p50 = values_sorted[len(values_sorted) // 2]
        p95 = values_sorted[min(len(values_sorted) - 1, int(len(values_sorted) * 0.95))]
        total = sum(values)
        total_by_stage[stage] = total
        lines.append(
            f"| {stage} | {len(values)} | {mean:.3f} | {p50:.3f} | {p95:.3f} | {total:.2f} |"
        )

    grand_total = sum(total_by_stage.values())
    extraction_total = total_by_stage.get("extract", 0.0) + total_by_stage.get(
        "resolve_and_load", 0.0
    )
    if grand_total > 0:
        extraction_share = extraction_total / grand_total * 100
        lines.append(
            f"\nExtraction+resolution stages account for **{extraction_share:.1f}%** of "
            "total measured pipeline time — "
            + (
                "confirms Phase 3's own claim that LLM extraction is the real bottleneck."
                if extraction_share > 50
                else "does NOT confirm Phase 3's claim at this sample size/content shape — "
                "extraction was not the dominant cost here."
            )
        )

        per_doc_seconds = grand_total / document_count
        lines.append(
            f"\nMeasured: ~{per_doc_seconds:.2f}s of total pipeline time per document "
            f"(all stages combined) at this batch size.\n"
        )
        lines.append("Extrapolated (not literally run, arithmetic only):\n")
        for n in (100, 1000, 10000):
            hours = per_doc_seconds * n / 3600
            lines.append(f"- At {n} documents: ~{hours:.1f} hours of total pipeline time")

    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--count", type=int, default=20)
    parser.add_argument("--per-doc-timeout", type=float, default=120.0)
    parser.add_argument("--api-url", type=str, default="http://localhost:8000")
    args = parser.parse_args()

    document_ids = []
    with httpx.Client(timeout=30.0) as client:
        for i in range(args.count):
            nonce = uuid.uuid4().hex[:8]
            data = _generate_docx(nonce)
            print(f"Uploading document {i + 1}/{args.count} ({nonce})...")
            document_id = _upload_and_wait(
                client, args.api_url, data, f"load_test_{nonce}.docx", args.per_doc_timeout
            )
            document_ids.append(document_id)

    db = SessionLocal()
    try:
        durations = _stage_durations(db, document_ids)
    finally:
        db.close()

    report = _report(durations, len(document_ids))
    print("\n" + report)

    output_path = (
        Path(__file__).resolve().parent.parent.parent / "docs" / "phase-8-load-test-results.md"
    )
    output_path.write_text(report, encoding="utf-8")
    print(f"Report written to {output_path}")


if __name__ == "__main__":
    main()
